"""Builds small upload fixtures in memory, so no binary files are committed."""
import io
import zipfile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter


def make_pdf(pages: list[str]) -> bytes:
    """A valid PDF with one line of text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(f'{3 + 2 * i} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>"]
    for i, text in enumerate(pages):
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                    "/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_blank_pdf(pages: int = 1) -> bytes:
    """PDF pages with no text at all (what a scanned document looks like to a text extractor)."""
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_encrypted_pdf(password: str = "secret") -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["Hidden text"]))))
    writer.encrypt(password, algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def make_zip(members: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return out.getvalue()


DOCX_TYPES = (b'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              b'<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument'
              b'.wordprocessingml.document.main+xml"/></Types>')
