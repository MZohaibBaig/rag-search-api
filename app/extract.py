"""Turn an uploaded file into clean plain text, or refuse it with a 4xx IngestError.

Parser libraries are imported inside each extractor so they only cost RAM when such a file arrives.
"""
import csv
import io
import re
import zipfile

from fastapi import HTTPException

MAX_TEXT_CHARS = 300_000
MAX_PDF_PAGES = 100
MAX_SHEETS = 20
MAX_ROWS = 2000                       # spreadsheet / csv rows scanned, header rows included
MAX_ZIP_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000

SUPPORTED = ".txt, .pdf, .docx, .xlsx or .csv"
_LEGACY = {".doc": ".docx", ".xls": ".xlsx"}
# docx/xlsx are zips; this text must appear in [Content_Types].xml for the extension to be believed
_OOXML_MARKER = {".docx": b"wordprocessingml", ".xlsx": b"spreadsheetml"}

_UNREADABLE = "Could not read this file"
_STRIP = re.compile(r"[\x00\ud800-\udfff]")     # NUL and lone surrogates (Postgres / UTF-8 reject them)
_INLINE_SPACE = re.compile(r"[^\S\n]+")
_BLANK_LINES = re.compile(r"\n{3,}")


class IngestError(HTTPException):
    """A refusal that is the client's to fix: raised anywhere in the upload path, returned as-is."""

    def __init__(self, status_code: int, detail: str):
        super().__init__(status_code=status_code, detail=detail)


def clean_text(text: str) -> str:
    """Drop NUL and invalid surrogates; collapse spaces; keep line and paragraph breaks."""
    text = _STRIP.sub("", text).replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(_INLINE_SPACE.sub(" ", line).strip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def clean_filename(filename: str | None) -> str:
    """Basename only (either slash style), no NUL/surrogates, at most 255 chars (extension kept
    so the type can still be told), never empty."""
    name = re.split(r"[\\/]", _STRIP.sub("", filename or ""))[-1].strip()
    ext = _extension(name)
    if len(name) > 255 and len(ext) < 255:
        name = name[:255 - len(ext)] + ext
    return name[:255] or "upload"


def _extension(filename: str) -> str:
    return filename[filename.rfind("."):].lower() if "." in filename else ""


def extract_text(filename: str, content: bytes) -> str:
    """Plain text of an uploaded file. Raises IngestError: 415 wrong type, 413 too big, 422 unreadable."""
    ext = _extension(clean_filename(filename))
    if ext in _LEGACY:
        raise IngestError(415, f"Please save as {_LEGACY[ext]}")
    if ext not in _EXTRACTORS:
        raise IngestError(415, f"Unsupported file type. Upload a {SUPPORTED} file.")
    _check_content(ext, content)

    try:
        text = _EXTRACTORS[ext](filename, content)
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(422, _UNREADABLE) from exc

    text = clean_text(text)
    if not text:
        raise IngestError(422, "No text found in this file")
    _check_length(len(text))
    return text


def _check_length(chars: int) -> None:
    if chars > MAX_TEXT_CHARS:
        raise IngestError(413, f"This file has too much text (limit {MAX_TEXT_CHARS:,} characters)")


def _check_content(ext: str, content: bytes) -> None:
    """The bytes must look like what the extension claims; for zip formats also bound the zip."""
    mismatch = IngestError(415, f"This file is not a valid {ext} file")
    if ext == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise mismatch
    elif ext in _OOXML_MARKER:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_ZIP_UNCOMPRESSED:
                    raise IngestError(413, "This file is too large when unpacked")
                if "[Content_Types].xml" not in archive.namelist():
                    raise mismatch
                if _OOXML_MARKER[ext] not in archive.read("[Content_Types].xml"):
                    raise mismatch
        except zipfile.BadZipFile as exc:
            raise mismatch from exc


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def _row_line(sheet: str, headers: list[str], values: list) -> str | None:
    """'Sheet | Header: value; Header: value.' (None for an empty row). The closing period keeps the
    chunker from mistaking a short row for a section heading."""
    pairs = [
        f"{headers[i] if i < len(headers) and headers[i] else f'Column {i + 1}'}: {str(v).strip()}"
        for i, v in enumerate(values) if v is not None and str(v).strip()
    ]
    return f"{sheet} | {'; '.join(pairs)}".rstrip(".") + "." if pairs else None


def _table_text(sheet: str, rows, state: dict) -> list[str]:
    """Row lines for one table; the first non-empty row is the header. `state` counts rows and
    characters across all tables of the file, so the limits apply to the file as a whole."""
    lines: list[str] = []
    headers: list[str] | None = None
    for row in rows:
        state["rows"] += 1
        if state["rows"] > MAX_ROWS:
            raise IngestError(413, f"This file has too many rows (limit {MAX_ROWS:,})")
        values = list(row)
        if headers is None:
            if any(v is not None and str(v).strip() for v in values):
                headers = [str(v).strip() if v is not None else "" for v in values]
            continue
        line = _row_line(sheet, headers, values)
        if line:
            lines.append(line)
            state["chars"] += len(line)
            _check_length(state["chars"])
    return lines


def _txt(filename: str, content: bytes) -> str:
    return _decode(content)


def _csv(filename: str, content: bytes) -> str:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    sheet = re.split(r"[\\/]", filename)[-1].rsplit(".", 1)[0]
    lines = _table_text(sheet, csv.reader(io.StringIO(text), dialect), {"rows": 0, "chars": 0})
    return "\n\n".join(lines)


def _pdf(filename: str, content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise IngestError(422, "This PDF is password-protected")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise IngestError(413, f"This PDF has too many pages (limit {MAX_PDF_PAGES})")
    pages: list[str] = []
    total = 0
    for page in reader.pages:           # one page at a time; stop as soon as the text limit is passed
        pages.append(page.extract_text() or "")
        total += len(pages[-1])
        _check_length(total)
    return "\n\n".join(pages)


def _docx(filename: str, content: bytes) -> str:
    from docx import Document
    from docx.table import Table

    blocks: list[str] = []
    for item in Document(io.BytesIO(content)).iter_inner_content():   # paragraphs and tables, in order
        if isinstance(item, Table):
            for row in item.rows:
                cells, previous = [], None
                for cell in row.cells:
                    if cell._tc is not previous:                      # merged cells repeat in row.cells
                        cells.append(cell.text.strip())
                    previous = cell._tc
                line = " | ".join(c for c in cells if c)
                if line:
                    blocks.append(line.rstrip(".") + ".")     # period: see _row_line
        else:
            blocks.append(item.text)
    return "\n\n".join(b for b in blocks if b.strip())


def _xlsx(filename: str, content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise IngestError(413, f"This workbook has too many sheets (limit {MAX_SHEETS})")
        state = {"rows": 0, "chars": 0}
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines += _table_text(sheet.title, sheet.iter_rows(values_only=True), state)
        return "\n\n".join(lines)
    finally:
        workbook.close()


_EXTRACTORS = {".txt": _txt, ".csv": _csv, ".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx}
